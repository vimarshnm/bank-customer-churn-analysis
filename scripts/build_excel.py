"""build_excel.py - build excel/churn_analysis.xlsx from data/extracts/*.csv.

The script places the data and writes Excel formulas as text.
Every KPI, segment, list and ROI figure in the workbook is an Excel formula
(COUNTIFS, AVERAGEIFS, SUMIFS, XLOOKUP, LET, FILTER, SORT, UNIQUE ...) that
Excel calculates when the file is opened - nothing is pre-computed here.

Usage: python3 scripts/build_excel.py [--force]   (after ./scripts/export_extracts.sh)

Safety: once the workbook has been saved in Excel it may contain manual edits,
so this script refuses to overwrite it. --force overwrites anyway.

Why the post-processing at the end: openpyxl cannot mark a formula as a
"dynamic array" (one that spills into neighbouring cells). Without that flag
Excel would show only the first value of FILTER/SORT/UNIQUE. We add the same
spill metadata Excel itself writes (xl/metadata.xml + cm="1" on the cell).
Requires Microsoft 365 or Excel 2021+.
"""
import csv
import re
import shutil
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
EXTRACTS = ROOT / "data" / "extracts"
OUT = ROOT / "excel" / "churn_analysis.xlsx"


def saved_by_excel(path):
    """True if the file was last saved by Excel rather than by this script."""
    try:
        with zipfile.ZipFile(path) as z:
            return b"Openpyxl" not in z.read("docProps/app.xml")
    except (KeyError, zipfile.BadZipFile):
        return True


if OUT.exists() and saved_by_excel(OUT) and "--force" not in sys.argv:
    sys.exit(f"REFUSING to overwrite {OUT.name}: it has been saved in Excel and may contain "
             "manual edits. Save a copy first, then "
             "rerun with --force.")

# ----------------------------------------------------------------------------
# Styles
# ----------------------------------------------------------------------------
FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=16, bold=True, color="1F4E78")
F_H2 = Font(name=FONT, size=12, bold=True, color="1F4E78")
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_INPUT = Font(name=FONT, size=10, bold=True, color="0000FF")   # hardcoded input
F_LINK = Font(name=FONT, size=10, color="008000")                # link to another sheet
FILL_HEAD = PatternFill("solid", fgColor="1F4E78")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_SUB = PatternFill("solid", fgColor="DDEBF7")
FILL_GREY = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FMT_INT = "#,##0"
FMT_PCT = "0.00%"
FMT_USD = "$#,##0"
FMT_USD2 = "$#,##0.00"
FMT_X = '0.00"x"'

# ----------------------------------------------------------------------------
# Formula helper: write formulas in normal Excel syntax, convert to file syntax
# ----------------------------------------------------------------------------
# Post-2010 functions are stored in the file with a prefix (Excel hides it).
PREFIX = {
    "XLOOKUP": "_xlfn.XLOOKUP", "XMATCH": "_xlfn.XMATCH", "LET": "_xlfn.LET",
    "FILTER": "_xlfn._xlws.FILTER", "SORT": "_xlfn._xlws.SORT", "UNIQUE": "_xlfn.UNIQUE",
    "SEQUENCE": "_xlfn.SEQUENCE", "SWITCH": "_xlfn.SWITCH",
}
DYNAMIC = {"XLOOKUP", "XMATCH", "LET", "FILTER", "SORT", "UNIQUE", "SEQUENCE"}
# A bare table name ("tblCustomers") does not resolve when stored in the file;
# Excel itself stores the whole data body as "tblCustomers[#Data]".
TABLES = {"tblCustomers", "tblKPIRef", "tblROIDefaults", "tblAgeBand", "tblBalanceBand",
          "tblTenureBand", "tblCreditBand", "tblCountry", "tblRiskTier", "tblCallPriority"}


def fx(formula, let_names=()):
    """Convert a readable formula to Excel's stored form.
    Returns (stored_formula, is_dynamic)."""
    out, i, dynamic = [], 0, False
    in_str, depth = False, 0
    while i < len(formula):
        ch = formula[i]
        if in_str:
            out.append(ch)
            if ch == '"':
                in_str = False
            i += 1
            continue
        if ch == '"':
            in_str = True
            out.append(ch)
            i += 1
            continue
        if ch == "[":
            depth += 1
        elif ch == "]":
            depth -= 1
        # sheet-qualified spill ref: 'At-Risk'!A10# or Summary!B5# -> ANCHORARRAY(sheet!ref)
        sheet_spill = re.match(r"('[^']+'|[A-Za-z_][A-Za-z0-9_.]*)!\$?[A-Z]{1,3}\$?\d+#", formula[i:])
        if depth == 0 and sheet_spill:
            out.append(f"_xlfn.ANCHORARRAY({sheet_spill.group(0)[:-1]})")
            dynamic = True
            i += len(sheet_spill.group(0))
            continue
        if depth == 0 and (ch.isalpha() or ch in "_$"):
            m = re.match(r"\$?[A-Za-z_][A-Za-z0-9_.]*\$?\d*#?", formula[i:])
            tok = m.group(0)
            nxt = formula[i + len(tok): i + len(tok) + 1]
            if tok.endswith("#") and re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+#", tok):
                out.append(f"_xlfn.ANCHORARRAY({tok[:-1]})")    # spill ref A10#
                dynamic = True
            elif nxt == "(" and tok.upper() in PREFIX:
                out.append(PREFIX[tok.upper()])
                dynamic = dynamic or tok.upper() in DYNAMIC
            elif tok in let_names:
                out.append("_xlpm." + tok)
            elif tok in TABLES and nxt != "[":
                out.append(tok + "[#Data]")
            else:
                out.append(tok)
            i += len(tok)
            continue
        out.append(ch)
        i += 1
    return "".join(out), dynamic


def put(ws, ref, formula, let_names=(), fmt=None, font=None):
    """Write a formula; dynamic-array formulas are flagged so they can spill."""
    stored, dynamic = fx(formula, let_names)
    cell = ws[ref]
    if dynamic:
        ws[ref] = ArrayFormula(ref, stored)
    else:
        cell.value = stored
    cell.font = font or F_BASE
    if fmt:
        cell.number_format = fmt
    return cell


def label(ws, ref, text, font=None, fill=None):
    c = ws[ref]
    c.value = text
    c.font = font or F_BASE
    if fill:
        c.fill = fill
    return c


def header_row(ws, row, col, names, widths=None):
    for j, name in enumerate(names):
        c = ws.cell(row=row, column=col + j, value=name)
        c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def read_csv(name):
    """Read an extract, converting numbers (int / float) and blanks (None)."""
    with open(EXTRACTS / f"{name}.csv", newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    head, body = rows[0], rows[1:]

    def conv(v):
        if v == "":
            return None
        for t in (int, float):
            try:
                return t(v)
            except ValueError:
                pass
        return v

    return head, [[conv(v) for v in r] for r in body]


def add_table(ws, name, top_left_row, top_left_col, head, body, style="TableStyleMedium2",
              formats=None):
    for j, h in enumerate(head):
        ws.cell(row=top_left_row, column=top_left_col + j, value=h).font = F_BOLD
    for i, r in enumerate(body, start=1):
        for j, v in enumerate(r):
            c = ws.cell(row=top_left_row + i, column=top_left_col + j, value=v)
            c.font = F_BASE
            if formats and head[j] in formats:
                c.number_format = formats[head[j]]
    ref = (f"{get_column_letter(top_left_col)}{top_left_row}:"
           f"{get_column_letter(top_left_col + len(head) - 1)}{top_left_row + len(body)}")
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True)
    ws.add_table(t)
    return ref


def define(wb, name, ref, comment=None):
    wb.defined_names[name] = DefinedName(name, attr_text=ref, comment=comment)


# ============================================================================
# Build
# ============================================================================
wb = Workbook()
wb.calculation.fullCalcOnLoad = True          # Excel recalculates on open

ws_guide = wb.active
ws_guide.title = "Guide"
ws_sum = wb.create_sheet("Summary")
ws_seg = wb.create_sheet("Segments")
ws_risk = wb.create_sheet("At-Risk")
ws_scn = wb.create_sheet("Scenario")
ws_data = wb.create_sheet("Data")
ws_lk = wb.create_sheet("Lookups")

# ----------------------------------------------------------------------------
# Data sheet: the customer-level fact table as an Excel Table
# ----------------------------------------------------------------------------
fact_head, fact_body = read_csv("fact_customers")
add_table(ws_data, "tblCustomers", 1, 1, fact_head, fact_body, style="TableStyleMedium2",
          formats={"balance": FMT_USD2, "estimated_salary": FMT_USD2})
ws_data.freeze_panes = "B2"
for j, h in enumerate(fact_head, start=1):
    ws_data.column_dimensions[get_column_letter(j)].width = max(10, min(len(h) + 3, 22))

# ----------------------------------------------------------------------------
# Lookups sheet: dimension tables, ROI defaults and SQL reference values
# ----------------------------------------------------------------------------
label(ws_lk, "A1", "Lookups - dimension tables and SQL reference values", F_TITLE)
label(ws_lk, "A2", "Loaded from data/extracts/*.csv (produced by sql/18_extract_views.sql). "
      "Do not edit: change the SQL and rebuild instead.", F_NOTE)
row = 4
lookup_tables = [
    ("kpi_reference", "tblKPIRef", "SQL answers for each headline KPI (Summary check column)"),
    ("roi_assumptions", "tblROIDefaults", "ROI default inputs from SQL (all are assumptions)"),
    ("dim_age_band", "tblAgeBand", "Age bands"),
    ("dim_balance_band", "tblBalanceBand", "Balance bands"),
    ("dim_tenure_band", "tblTenureBand", "Tenure bands"),
    ("dim_credit_score_band", "tblCreditBand", "Credit score bands"),
    ("dim_country", "tblCountry", "Countries"),
    ("dim_risk_tier", "tblRiskTier", "Risk tiers"),
    ("dim_call_priority", "tblCallPriority", "Call priority (top 10% only)"),
]
dims = {}
for csv_name, tbl, title in lookup_tables:
    head, body = read_csv(csv_name)
    dims[csv_name] = (head, body)
    label(ws_lk, f"A{row}", title, F_H2)
    add_table(ws_lk, tbl, row + 1, 1, head, body, style="TableStyleLight9")
    row += len(body) + 4

# Country dropdown options (used by data validation on Summary and At-Risk)
label(ws_lk, f"A{row}", "Country dropdown options", F_H2)
label(ws_lk, f"A{row + 1}", "country_option", F_BOLD)
options = ["All"] + [r[0] for r in dims["dim_country"][1]]
for k, opt in enumerate(options, start=row + 2):
    label(ws_lk, f"A{k}", opt)
define(wb, "CountryOptions", f"Lookups!$A${row + 2}:$A${row + 1 + len(options)}")
set_widths(ws_lk, {"A": 30, "B": 18, "C": 16, "D": 22, "E": 14, "F": 48})

# ----------------------------------------------------------------------------
# Scenario sheet: ROI model driven by four input cells
# (built before Summary because Summary links to its results)
# ----------------------------------------------------------------------------
s = ws_scn
label(s, "A1", "Scenario - retention campaign cost-benefit", F_TITLE)
label(s, "A2", "Change the yellow input cells; every result below updates. "
      "All inputs are assumptions; replace them with the bank's real figures.", F_NOTE)
header_row(s, 4, 1, ["Input", "Value", "SQL default", "Unit", "Note"])
roi_rows = [
    ("Cost per call", "CostPerCall", "cost_per_call", 10, FMT_USD2, "USD per customer called",
     "Assumption: agent time + admin. Every customer on the list is called."),
    ("Save rate", "SaveRate", "save_rate", 0.20, "0.0%", "share of would-be leavers saved",
     "Assumption: share of called would-be leavers who stay because of the call."),
    ("Deposit margin", "DepositMargin", "deposit_margin_rate", 0.01, "0.00%", "of balance per year",
     "Assumption: annual profit on a retained balance (net interest margin)."),
    ("Fixed value per customer", "FixedValue", "value_per_customer_fixed", 100, FMT_USD2,
     "USD per retained customer per year", "Assumption: fees and other products, any balance."),
]
for k, (name, rng, param, default, fmt, unit, note) in enumerate(roi_rows, start=5):
    label(s, f"A{k}", name, F_BOLD)
    c = s[f"B{k}"]
    c.value, c.font, c.fill, c.number_format, c.border = default, F_INPUT, FILL_INPUT, fmt, BOX
    c.comment = Comment(f"{note} Default {default} comes from SQL table roi_assumptions.", "Analyst")
    put(s, f"C{k}", f'=XLOOKUP("{param}", tblROIDefaults[parameter], tblROIDefaults[value], "missing")',
        fmt=fmt, font=F_LINK)
    label(s, f"D{k}", unit)
    label(s, f"E{k}", note, F_NOTE)
    define(wb, rng, f"Scenario!$B${k}")
label(s, "A10", "Inputs at SQL defaults?", F_BOLD)
put(s, "B10", "=AND(B5=C5, B6=C6, B7=C7, B8=C8)")
define(wb, "InputsAtDefault", "Scenario!$B$10")
label(s, "C10", "When FALSE, Summary marks the ROI checks 'n/a - inputs changed'.", F_NOTE)

header_row(s, 13, 1, ["Metric", "Top 10%", "Top 20%", "Decile 2 only (next 10%)", "How it is calculated"])
scen = [
    # (label, top10 formula, top20 formula, decile-2 formula, format, explanation)
    ("Customers called",
     "=COUNTIFS(tblCustomers[in_top10], 1)", "=COUNTIFS(tblCustomers[in_top20], 1)", "=C14-B14",
     FMT_INT, "COUNTIFS on the in_top10 / in_top20 flag"),
    ("Would-be leavers on the list",
     "=COUNTIFS(tblCustomers[in_top10], 1, tblCustomers[exited], 1)",
     "=COUNTIFS(tblCustomers[in_top20], 1, tblCustomers[exited], 1)", "=C15-B15",
     FMT_INT, "Listed customers who actually left (stand-in for who will leave)"),
    ("Deposits of would-be leavers",
     "=SUMIFS(tblCustomers[balance], tblCustomers[in_top10], 1, tblCustomers[exited], 1)",
     "=SUMIFS(tblCustomers[balance], tblCustomers[in_top20], 1, tblCustomers[exited], 1)", "=C16-B16",
     FMT_USD, "SUMIFS of balance for listed leavers"),
    ("Customers saved", "=B15*SaveRate", "=C15*SaveRate", "=D15*SaveRate",
     "#,##0.0", "Would-be leavers x SaveRate"),
    ("Deposits retained", "=B16*SaveRate", "=C16*SaveRate", "=D16*SaveRate",
     FMT_USD, "Leaver deposits x SaveRate"),
    ("Campaign cost", "=B14*CostPerCall", "=C14*CostPerCall", "=D14*CostPerCall",
     FMT_USD2, "Customers called x CostPerCall"),
    ("Value if every would-be leaver saved",
     "=DepositMargin*B16 + FixedValue*B15", "=DepositMargin*C16 + FixedValue*C15",
     "=DepositMargin*D16 + FixedValue*D15",
     FMT_USD2, "DepositMargin x leaver deposits + FixedValue x leavers (1 year)"),
    ("Value saved", "=SaveRate*B20", "=SaveRate*C20", "=SaveRate*D20",
     FMT_USD2, "SaveRate x value if all saved"),
    ("Net benefit", "=B21-B19", "=C21-C19", "=D21-D19", FMT_USD2, "Value saved - campaign cost"),
    ("ROI", "=IFERROR(B22/B19, 0)", "=IFERROR(C22/C19, 0)", "=IFERROR(D22/D19, 0)",
     FMT_X, "Net benefit / campaign cost"),
    ("Break-even save rate", "=IFERROR(B19/B20, 0)", "=IFERROR(C19/C20, 0)", "=IFERROR(D19/D20, 0)",
     FMT_PCT, "Save rate at which net benefit = 0 (cost / value if all saved)"),
    ("Break-even cost per call", "=IFERROR(B21/B14, 0)", "=IFERROR(C21/C14, 0)",
     "=IFERROR(D21/D14, 0)", FMT_USD2, "Call cost at which net benefit = 0 (value saved / calls)"),
]
for k, (name, f10, f20, fd2, fmt, how) in enumerate(scen, start=14):
    label(s, f"A{k}", name, F_BOLD)
    put(s, f"B{k}", f10, fmt=fmt)
    put(s, f"C{k}", f20, fmt=fmt)
    put(s, f"D{k}", fd2, fmt=fmt)
    label(s, f"E{k}", how, F_NOTE)
    for col in "ABCD":
        s[f"{col}{k}"].border = BOX
for col in "ABCD":
    s[f"{col}22"].fill = FILL_SUB
    s[f"{col}22"].font = F_BOLD

set_widths(s, {"A": 36, "B": 18, "C": 18, "D": 24, "E": 62})
s.freeze_panes = "A4"

# ----------------------------------------------------------------------------
# Summary sheet: live KPIs + reconciliation against SQL, country selector
# ----------------------------------------------------------------------------
s = ws_sum
label(s, "A1", "Summary - headline KPIs", F_TITLE)
label(s, "A2", "Excel value = live formula on the Data sheet. SQL value = XLOOKUP into "
      "Lookups!tblKPIRef (exported from SQL). Check = Excel rounded to the SQL precision.", F_NOTE)
header_row(s, 4, 1, ["KPI", "KPI code", "Excel value", "SQL value", "Difference", "Check",
                     "How the Excel value is calculated"])
T = "tblCustomers"
kpis = [
    ("Total customers", "total_customers", f"=COUNTA({T}[customer_id])", FMT_INT,
     "COUNTA of customer_id"),
    ("Leavers", "leavers", f"=COUNTIFS({T}[exited], 1)", FMT_INT, "COUNTIFS exited = 1"),
    ("Churn rate", "churn_rate", "=C6/C5", FMT_PCT, "Leavers / customers"),
    ("Churn rate - Germany", "churn_rate_germany",
     f'=AVERAGEIFS({T}[exited], {T}[country], "Germany")', FMT_PCT,
     "AVERAGEIFS of the 0/1 exited flag = churn rate"),
    ("Churn rate - inactive members", "churn_rate_inactive",
     f"=AVERAGEIFS({T}[exited], {T}[is_active_member], 0)", FMT_PCT, "AVERAGEIFS, inactive only"),
    ("Churn rate - age 50-59", "churn_rate_age_50_59",
     f'=AVERAGEIFS({T}[exited], {T}[age_band], "50-59")', FMT_PCT, "AVERAGEIFS, age band 50-59"),
    ("Total deposits", "total_deposits", f"=SUM({T}[balance])", FMT_USD, "SUM of balance"),
    ("Deposits lost (leavers' balances)", "deposits_lost",
     f"=SUMIFS({T}[balance], {T}[exited], 1)", FMT_USD, "SUMIFS balance where exited = 1"),
    ("Top 10% - churn rate", "top10_churn_rate",
     f"=AVERAGEIFS({T}[exited], {T}[in_top10], 1)", FMT_PCT, "AVERAGEIFS, riskiest 10%"),
    ("Top 10% - lift", "top10_lift", "=C13/C7", FMT_X, "Top-10% churn / overall churn"),
    ("Top 10% - share of leavers captured", "top10_capture",
     f"=COUNTIFS({T}[in_top10], 1, {T}[exited], 1) / C6", FMT_PCT,
     "Leavers in top 10% / all leavers"),
    ("Top 10% - deposits held", "top10_deposits_held",
     f"=SUMIFS({T}[balance], {T}[in_top10], 1)", FMT_USD, "SUMIFS balance, riskiest 10%"),
    ("Top 10% campaign - net benefit", "top10_campaign_net_benefit", "=Scenario!B22", FMT_USD2,
     "Link to Scenario sheet (depends on inputs)"),
    ("Top 10% campaign - break-even save rate", "top10_break_even_save_rate", "=Scenario!B24",
     FMT_PCT, "Link to Scenario sheet"),
    ("Top 10% campaign - ROI", "top10_roi", "=Scenario!B23", FMT_X, "Link to Scenario sheet"),
    ("Top 20% campaign - net benefit", "top20_campaign_net_benefit", "=Scenario!C22", FMT_USD2,
     "Link to Scenario sheet"),
    ("Top 20% campaign - ROI", "top20_roi", "=Scenario!C23", FMT_X, "Link to Scenario sheet"),
    ("Top 20% campaign - break-even save rate", "top20_break_even_save_rate", "=Scenario!C24",
     FMT_PCT, "Link to Scenario sheet"),
]
first = 5
for k, (name, code, formula, fmt, how) in enumerate(kpis, start=first):
    label(s, f"A{k}", name, F_BOLD)
    label(s, f"B{k}", code, F_NOTE)
    put(s, f"C{k}", formula, fmt=fmt, font=F_LINK if "Scenario!" in formula else None)
    put(s, f"D{k}", f'=XLOOKUP(B{k}, tblKPIRef[kpi], tblKPIRef[sql_value], "missing")',
        fmt=fmt, font=F_LINK)
    put(s, f"E{k}", f"=IFERROR(C{k}-D{k}, \"\")", fmt=fmt)
    check = (f'=LET(fmt, XLOOKUP(B{k}, tblKPIRef[kpi], tblKPIRef[format]), '
             f'dp, SWITCH(fmt, "count", 0, "percent", 4, "ratio", 2, "currency", 2), '
             f'IF(ABS(ROUND(C{k}, dp) - D{k}) < 0.000001, "Match", "Mismatch"))')
    if "Scenario!" in formula:
        check = '=IF(NOT(InputsAtDefault), "n/a - inputs changed", ' + check[1:] + ")"
    put(s, f"F{k}", check, let_names=("fmt", "dp"))
    label(s, f"G{k}", how, F_NOTE)
    for col in "ABCDEF":
        s[f"{col}{k}"].border = BOX
last = first + len(kpis) - 1
rng = f"F{first}:F{last}"
s.conditional_formatting.add(rng, FormulaRule(formula=[f'F{first}="Match"'],
                                              fill=PatternFill("solid", fgColor="C6EFCE")))
s.conditional_formatting.add(rng, FormulaRule(formula=[f'F{first}="Mismatch"'],
                                              fill=PatternFill("solid", fgColor="FFC7CE")))
label(s, f"A{last + 1}", "All checks", F_BOLD)
put(s, f"F{last + 1}",
    f'=IF(COUNTIF({rng}, "Mismatch") = 0, "All match SQL", COUNTIF({rng}, "Mismatch") & " mismatch(es)")',
    font=F_BOLD)

# Country selector -----------------------------------------------------------
top = last + 4                                              # row 26
label(s, f"A{top}", "Country view", F_H2)
label(s, f"A{top + 1}", "Select country ▸", F_BOLD)
c = s[f"B{top + 1}"]
c.value, c.font, c.fill, c.border = "All", F_INPUT, FILL_INPUT, BOX
c.comment = Comment("Dropdown: pick All, France, Germany or Spain.", "Analyst")
define(wb, "SelectedCountry", f"Summary!$B${top + 1}")
# formula1 is stored WITHOUT a leading "=" in the file format
dv = DataValidation(type="list", formula1="CountryOptions", allow_blank=False,
                    showErrorMessage=True)
dv.error, dv.errorTitle = "Pick a country from the list.", "Invalid country"
s.add_data_validation(dv)
dv.add(f"B{top + 1}")
label(s, f"C{top + 1}", "Dropdown (data validation). All KPIs below follow it.", F_NOTE)

header_row(s, top + 3, 1, ["Measure", "Value", "How it is calculated"])
crit = 'LET(c, IF(SelectedCountry="All", "*", SelectedCountry), '   # "*" = any country
r0 = top + 4
country_kpis = [
    ("Country code", '=IF(SelectedCountry="All", "All", XLOOKUP(SelectedCountry, tblCountry[country], tblCountry[iso_code], "?"))',
     None, "XLOOKUP into the country dimension"),
    ("Customers", f"={crit}COUNTIFS({T}[country], c))", FMT_INT, 'COUNTIFS; "All" uses the * wildcard'),
    ("Leavers", f"={crit}COUNTIFS({T}[country], c, {T}[exited], 1))", FMT_INT, "COUNTIFS exited = 1"),
    ("Churn rate", f"=IFERROR(B{r0 + 2}/B{r0 + 1}, 0)", FMT_PCT, "Leavers / customers"),
    ("Lift vs whole bank", f"=IFERROR(B{r0 + 3}/C7, 0)", FMT_X, "Country churn / overall churn"),
    ("Share of all leavers", f"=IFERROR(B{r0 + 2}/C6, 0)", FMT_PCT, "Country leavers / all leavers"),
    ("Average balance", f"={crit}AVERAGEIFS({T}[balance], {T}[country], c))", FMT_USD2, "AVERAGEIFS"),
    ("Deposits held", f"={crit}SUMIFS({T}[balance], {T}[country], c))", FMT_USD, "SUMIFS"),
    ("Deposits lost", f"={crit}SUMIFS({T}[balance], {T}[country], c, {T}[exited], 1))", FMT_USD,
     "SUMIFS where exited = 1"),
    ("Share of country deposits lost", f"=IFERROR(B{r0 + 8}/B{r0 + 7}, 0)", FMT_PCT, "Lost / held"),
    ("Customers in the top 10%", f"={crit}COUNTIFS({T}[country], c, {T}[in_top10], 1))", FMT_INT,
     "COUNTIFS in_top10 = 1"),
    ("Top-10% deposits held", f"={crit}SUMIFS({T}[balance], {T}[country], c, {T}[in_top10], 1))",
     FMT_USD, "SUMIFS in_top10 = 1"),
]
for k, (name, formula, fmt, how) in enumerate(country_kpis, start=r0):
    label(s, f"A{k}", name, F_BOLD)
    put(s, f"B{k}", formula, let_names=("c",), fmt=fmt)
    label(s, f"C{k}", how, F_NOTE)
    for col in "AB":
        s[f"{col}{k}"].border = BOX
set_widths(s, {"A": 40, "B": 26, "C": 20, "D": 20, "E": 14, "F": 22, "G": 46})
s.freeze_panes = "A5"

# ----------------------------------------------------------------------------
# Segments sheet: churn breakdowns with COUNTIFS / AVERAGEIFS
# ----------------------------------------------------------------------------
s = ws_seg
label(s, "A1", "Segments - churn by customer attribute", F_TITLE)
label(s, "A2", "Every number is a COUNTIFS / SUMIFS formula on the Data sheet. "
      "Matches SQL file sql/08_churn_by_attribute.sql (Q8.2) and sql/09 (Q9.1).", F_NOTE)
blocks = [
    ("Country", "country", [r[0] for r in dims["dim_country"][1]]),
    ("Gender", "gender", ["Female", "Male"]),
    ("Age band", "age_band", [r[0] for r in dims["dim_age_band"][1]]),
    ("Products held", "num_products", [1, 2, 3, 4]),
    ("Activity status", "activity_status", ["Active", "Inactive"]),
    ("Balance band", "balance_band", [r[0] for r in dims["dim_balance_band"][1]]),
    ("Tenure band", "tenure_band", [r[0] for r in dims["dim_tenure_band"][1]]),
    ("Credit score band", "credit_score_band", [r[0] for r in dims["dim_credit_score_band"][1]]),
    ("Has credit card (1 = yes)", "has_credit_card", [1, 0]),
    ("Risk tier", "risk_tier", [r[0] for r in dims["dim_risk_tier"][1]]),
]
row = 4
for title, col, values in blocks:
    label(s, f"A{row}", title, F_H2)
    header_row(s, row + 1, 1, [title, "Customers", "Leavers", "Churn rate",
                               "Share of all leavers", "Lift vs average"])
    start = row + 2
    for k, v in enumerate(values, start=start):
        lab = label(s, f"A{k}", v, F_BOLD)
        lab.alignment = Alignment(horizontal="left")
        put(s, f"B{k}", f"=COUNTIFS({T}[{col}], $A{k})", fmt=FMT_INT)
        put(s, f"C{k}", f"=COUNTIFS({T}[{col}], $A{k}, {T}[exited], 1)", fmt=FMT_INT)
        put(s, f"D{k}", f"=IFERROR(C{k}/B{k}, 0)", fmt=FMT_PCT)
        put(s, f"E{k}", f"=IFERROR(C{k}/COUNTIFS({T}[exited], 1), 0)", fmt=FMT_PCT)
        put(s, f"F{k}", f"=IFERROR(D{k}/AVERAGE({T}[exited]), 0)", fmt=FMT_X)
        for cc in "ABCDEF":
            s[f"{cc}{k}"].border = BOX
    end = start + len(values) - 1
    tot = end + 1
    label(s, f"A{tot}", "Total (should be 10,000 / 2,038)", F_NOTE).fill = FILL_GREY
    put(s, f"B{tot}", f"=SUM(B{start}:B{end})", fmt=FMT_INT).fill = FILL_GREY
    put(s, f"C{tot}", f"=SUM(C{start}:C{end})", fmt=FMT_INT).fill = FILL_GREY
    put(s, f"D{tot}", f"=IFERROR(C{tot}/B{tot}, 0)", fmt=FMT_PCT).fill = FILL_GREY
    s.conditional_formatting.add(
        f"D{start}:D{end}",
        ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                       end_type="num", end_value=1, end_color="F8696B"))
    row = tot + 3

# Age band x activity heatmap (feeds the same view as the Tableau heatmap)
label(s, "H4", "Age band x activity - churn rate", F_H2)
header_row(s, 5, 8, ["Age band", "Active", "Inactive"])
label(s, "L4", "Customers", F_H2)
header_row(s, 5, 12, ["Active", "Inactive"])
for k, band in enumerate([r[0] for r in dims["dim_age_band"][1]], start=6):
    label(s, f"H{k}", band, F_BOLD)
    for colL, status in (("I", "Active"), ("J", "Inactive")):
        put(s, f"{colL}{k}",
            f'=IFERROR(AVERAGEIFS({T}[exited], {T}[age_band], $H{k}, {T}[activity_status], {colL}$5), 0)',
            fmt=FMT_PCT)
        s[f"{colL}{k}"].border = BOX
    for colL, hdr in (("L", "I"), ("M", "J")):
        put(s, f"{colL}{k}",
            f"=COUNTIFS({T}[age_band], $H{k}, {T}[activity_status], {hdr}$5)", fmt=FMT_INT)
        s[f"{colL}{k}"].border = BOX
s.conditional_formatting.add(
    "I6:J11", ColorScaleRule(start_type="num", start_value=0, start_color="63BE7B",
                             mid_type="num", mid_value=0.2038, mid_color="FFEB84",
                             end_type="num", end_value=1, end_color="F8696B"))
label(s, "H13", "Colour: green = below the 20.38% average, red = far above. "
      "Small groups: 60-69 Inactive (94) and 70+ Inactive (17).", F_NOTE)
set_widths(s, {"A": 30, "B": 12, "C": 11, "D": 12, "E": 14, "F": 12, "G": 3,
               "H": 12, "I": 12, "J": 12, "K": 3, "L": 11, "M": 11})

# ----------------------------------------------------------------------------
# At-Risk sheet: the top-10% call list with dynamic arrays
# ----------------------------------------------------------------------------
s = ws_risk
label(s, "A1", "At-Risk - monthly call list (riskiest 10%)", F_TITLE)
label(s, "A2", "Built with one LET + FILTER + SORT formula in cell A10 that spills down. "
      "Change the country to re-filter; the list re-sorts by risk rank.", F_NOTE)
label(s, "A4", "Show country ▸", F_BOLD)
c = s["B4"]
c.value, c.font, c.fill, c.border = "All", F_INPUT, FILL_INPUT, BOX
define(wb, "AtRiskCountry", "'At-Risk'!$B$4")
dv2 = DataValidation(type="list", formula1="CountryOptions", allow_blank=False,
                     showErrorMessage=True)
s.add_data_validation(dv2)
dv2.add("B4")
label(s, "A5", "Customers listed", F_BOLD)
put(s, "B5", "=IFERROR(ROWS(A10#), 0)", fmt=FMT_INT)
label(s, "A6", "Deposits on the list", F_BOLD)
put(s, "B6", "=IFERROR(SUM(INDEX(A10#, 0, 8)), 0)", fmt=FMT_USD)

# Row 9 holds real Data column names: the formula looks them up with XMATCH,
# so the list shows whichever columns are named here.
list_cols = ["risk_rank", "customer_id", "surname", "country", "age", "activity_status",
             "num_products", "balance", "risk_score", "call_priority", "churn_status"]
friendly = ["Rank", "Customer ID", "Surname", "Country", "Age", "Activity", "Products",
            "Balance", "Risk score", "Call priority", "Outcome (snapshot)"]
for j, f in enumerate(friendly, start=1):
    c = s.cell(row=8, column=j, value=f)
    c.font, c.fill, c.border = F_HEAD, FILL_HEAD, BOX
    c.alignment = Alignment(horizontal="center", wrap_text=True)
for j, h in enumerate(list_cols, start=1):
    c = s.cell(row=9, column=j, value=h)
    c.font, c.fill = F_NOTE, FILL_SUB
label(s, "A7", "Row 9 = Data column names used by the formula (edit a name to swap a column).", F_NOTE)
at_risk = (
    "=LET("
    "picked, INDEX(tblCustomers, SEQUENCE(ROWS(tblCustomers)), XMATCH($A$9:$K$9, tblCustomers[#Headers])), "
    "keep, (tblCustomers[risk_decile] = 1) * IF(AtRiskCountry = \"All\", 1, tblCustomers[country] = AtRiskCountry), "
    "SORT(FILTER(picked, keep, \"No customers\"), 1, 1))"
)
put(s, "A10", at_risk, let_names=("picked", "keep"))
MAXR = 1009                                                  # top 10% = 1,000 rows max
s["H10"].number_format = FMT_USD2
for r in range(10, MAXR + 1):
    s[f"H{r}"].number_format = FMT_USD2
s.conditional_formatting.add(
    f"I10:I{MAXR}", ColorScaleRule(start_type="min", start_color="FFEB84",
                                   end_type="max", end_color="F8696B"))
s.conditional_formatting.add(f"H10:H{MAXR}", DataBarRule(start_type="num", start_value=0,
                                                         end_type="max", color="5B9BD5"))
for text, colour in (("P1", "F8CBAD"), ("P2", "FCE4D6"), ("P3", "FFF2CC")):
    s.conditional_formatting.add(
        f"J10:J{MAXR}", FormulaRule(formula=[f'LEFT($J10,2)="{text}"'],
                                    fill=PatternFill("solid", fgColor=colour)))

# Priority summary with UNIQUE (spills) + COUNTIFS/SUMIFS on the spilled range
label(s, "M4", "Call list by priority", F_H2)
header_row(s, 5, 13, ["Call priority", "Customers", "Would-be leavers", "Deposits held"])
crit_r = 'LET(c, IF(AtRiskCountry="All", "*", AtRiskCountry), '
put(s, "M6", '=SORT(UNIQUE(FILTER(tblCustomers[call_priority], (tblCustomers[risk_decile] = 1) * '
             'IF(AtRiskCountry = "All", 1, tblCustomers[country] = AtRiskCountry), "None")))')
put(s, "N6", f"={crit_r}COUNTIFS({T}[call_priority], M6#, {T}[risk_decile], 1, {T}[country], c))",
    let_names=("c",), fmt=FMT_INT)
put(s, "O6", f"={crit_r}COUNTIFS({T}[call_priority], M6#, {T}[risk_decile], 1, {T}[country], c, "
             f"{T}[exited], 1))", let_names=("c",), fmt=FMT_INT)
put(s, "P6", f"={crit_r}SUMIFS({T}[balance], {T}[call_priority], M6#, {T}[risk_decile], 1, "
             f"{T}[country], c))", let_names=("c",), fmt=FMT_USD)
for r in range(6, 10):
    s[f"N{r}"].number_format = FMT_INT
    s[f"O{r}"].number_format = FMT_INT
    s[f"P{r}"].number_format = FMT_USD
label(s, "M11", "M6 = SORT(UNIQUE(FILTER(...))) spills the priorities present; "
      "N6:P6 use M6# (the whole spilled range) as COUNTIFS / SUMIFS criteria.", F_NOTE)
set_widths(s, {"A": 8, "B": 13, "C": 16, "D": 10, "E": 6, "F": 10, "G": 9, "H": 15, "I": 10,
               "J": 15, "K": 12, "L": 3, "M": 16, "N": 11, "O": 15, "P": 16})
s.freeze_panes = "A10"

# ----------------------------------------------------------------------------
# Guide sheet
# ----------------------------------------------------------------------------
s = ws_guide
label(s, "A1", "Bank Customer Churn Analysis - Excel workbook", F_TITLE)
lines = [
    ("Purpose", "Live KPIs, churn segments, the monthly at-risk call list and a campaign ROI model "
                "for 10,000 bank customers. Every figure is an Excel formula on the Data sheet."),
    ("Source", "data/extracts/*.csv, produced by the SQL pipeline (sql/18_extract_views.sql). "
               "Rebuild: ./scripts/export_extracts.sh then python3 scripts/build_excel.py."),
    ("Requires", "Microsoft 365 or Excel 2021+ (dynamic arrays: LET, FILTER, SORT, UNIQUE, "
                 "XLOOKUP, XMATCH, SEQUENCE)."),
]
for k, (a, b) in enumerate(lines, start=3):
    label(s, f"A{k}", a, F_BOLD)
    label(s, f"B{k}", b).alignment = Alignment(wrap_text=True, vertical="top")
header_row(s, 7, 1, ["Sheet", "What it shows"])
sheets = [
    ("Summary", "18 headline KPIs with a check column against SQL; country dropdown view."),
    ("Segments", "Churn rate, share of leavers and lift by 10 attributes; age x activity heatmap."),
    ("At-Risk", "Top-10% call list (LET + FILTER + SORT), priority summary (UNIQUE), country filter."),
    ("Scenario", "Campaign ROI model: 4 input cells drive cost, value saved, net benefit, ROI, break-even."),
    ("Data", "tblCustomers: one row per customer (36 columns, see docs/data_dictionary.md)."),
    ("Lookups", "Dimension tables, ROI defaults and tblKPIRef (the SQL answers)."),
]
for k, (a, b) in enumerate(sheets, start=8):
    label(s, f"A{k}", a, F_BOLD)
    label(s, f"B{k}", b)
header_row(s, 15, 1, ["Legend", "Meaning"])
legend = [
    ("Input", "Blue bold on yellow = a value you may change (dropdowns, ROI inputs).", F_INPUT, FILL_INPUT),
    ("Formula", "Black = calculated by a formula. Do not overwrite.", F_BASE, None),
    ("Link", "Green = pulls a value from another sheet or lookup table.", F_LINK, None),
    ("Match", "Check column: Excel value equals the SQL value at SQL precision.", F_BASE,
     PatternFill("solid", fgColor="C6EFCE")),
]
for k, (a, b, f, fill) in enumerate(legend, start=16):
    c = label(s, f"A{k}", a, f, fill)
    label(s, f"B{k}", b)
header_row(s, 21, 1, ["Named range", "Refers to"])
names = [("SelectedCountry", "Summary!B27 - country dropdown"),
         ("AtRiskCountry", "At-Risk!B4 - call-list country filter"),
         ("CostPerCall / SaveRate / DepositMargin / FixedValue", "Scenario!B5:B8 - ROI inputs"),
         ("InputsAtDefault", "Scenario!B10 - TRUE when inputs equal the SQL defaults"),
         ("CountryOptions", "Lookups - list behind both dropdowns")]
for k, (a, b) in enumerate(names, start=22):
    label(s, f"A{k}", a, F_BOLD)
    label(s, f"B{k}", b)
label(s, "A28", "Example: on Summary, change 'Select country' to Germany - churn rate should "
      "read 32.44% and Customers 2,509.", F_NOTE)
set_widths(s, {"A": 48, "B": 100})

# sanity: the Guide's cell references must match the layout above
assert top + 1 == 27, "Guide text says Summary!B27; update it if the layout moves"

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = ws.title in ("Data", "Lookups")

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)

# ----------------------------------------------------------------------------
# Post-process: flag array formulas as dynamic arrays (spill), as Excel does
# ----------------------------------------------------------------------------
METADATA = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
    'xmlns:xda="http://schemas.microsoft.com/office/spreadsheetml/2017/dynamicarray">'
    '<metadataTypes count="1"><metadataType name="XLDAPR" minSupportedVersion="120000" copy="1" '
    'pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" '
    'clearComments="1" assign="1" coerce="1" cellMeta="1"/></metadataTypes>'
    '<futureMetadata name="XLDAPR" count="1"><bk><extLst><ext uri="{bdbb8cdc-fa1e-496e-a857-3c3f30c029c3}">'
    '<xda:dynamicArrayProperties fDynamic="1" fCollapsed="0"/></ext></extLst></bk></futureMetadata>'
    '<cellMetadata count="1"><bk><rc t="1" v="0"/></bk></cellMetadata></metadata>'
)
tmp = OUT.with_suffix(".tmp")
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.startswith("xl/worksheets/sheet"):
            data = re.sub(rb'<c r="([A-Z]+\d+)"([^>]*)><f t="array"',
                          rb'<c r="\1"\2 cm="1"><f t="array"', data)
        elif item.filename == "[Content_Types].xml":
            data = data.replace(b"</Types>", b'<Override PartName="/xl/metadata.xml" ContentType='
                                b'"application/vnd.openxmlformats-officedocument.spreadsheetml.'
                                b'sheetMetadata+xml"/></Types>')
        elif item.filename == "xl/_rels/workbook.xml.rels":
            data = data.replace(b"</Relationships>", b'<Relationship Id="rIdMeta1" Type='
                                b'"http://schemas.openxmlformats.org/officeDocument/2006/'
                                b'relationships/sheetMetadata" Target="metadata.xml"/></Relationships>')
        zout.writestr(item, data)
    zout.writestr("xl/metadata.xml", METADATA)
shutil.move(tmp, OUT)
print(f"Built {OUT.relative_to(ROOT)}")
